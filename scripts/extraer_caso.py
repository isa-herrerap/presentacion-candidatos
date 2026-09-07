#!/usr/bin/env python3
"""Extrae a texto plano los documentos de un caso: el perfil de cargo (.docx),
la planilla de screening (.xlsx) y, si la consultora lo dejó, el docx de
universo del proceso.

Uso:
    python3 scripts/extraer_caso.py <carpeta-del-caso>

Deja los resultados en <carpeta-del-caso>/_extraccion/:
    perfil.txt     — párrafos + tablas del docx del perfil ("campo | valor")
    screening.txt  — todas las hojas del xlsx ("col1 | col2 | ..."), incluyendo
                     comentarios de celda si los hay
    universo.txt   — SOLO si hay un .docx cuyo nombre de archivo contiene
                     "universo" (sin distinguir mayúsculas) — el total de
                     candidatos, el reparto hunting/postulación y los tres
                     repartos de LinkedIn del grupo hunting, que ningún otro
                     documento trae. Si existe, la skill generar-deck lee estos
                     datos de acá en vez de preguntarlos.

Los CV en PDF no se extraen acá: Claude los lee directo con su Read tool. El
perfil de cargo a veces también llega en PDF en vez de .docx — ese caso
tampoco se extrae acá (este script sólo sabe leer .docx): se lee igual que un
CV, directo con el Read tool. El script avisa cuando no encuentra ningún
.docx de perfil, para que no se te pase por alto.

Dependencias: openpyxl y python-docx. Si no están instaladas, el script
intenta usar el venv del proyecto (.venv/) y, si tampoco existe, imprime
las instrucciones exactas para crearlo.
"""
import glob
import os
import subprocess
import sys

PROYECTO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
VENV_PY = os.path.join(PROYECTO, ".venv", "bin", "python3")

try:
    import openpyxl
    from docx import Document
except ImportError:
    # Reintentar con el venv del proyecto, si existe y no somos ya él.
    if os.path.exists(VENV_PY) and os.path.abspath(sys.executable) != os.path.abspath(VENV_PY):
        os.execv(VENV_PY, [VENV_PY] + sys.argv)
    sys.exit(
        "Faltan dependencias (openpyxl, python-docx). Créalas una sola vez con:\n"
        f"  python3 -m venv {os.path.join(PROYECTO, '.venv')}\n"
        f"  {os.path.join(PROYECTO, '.venv', 'bin', 'pip')} install openpyxl python-docx\n"
        "y vuelve a correr este script."
    )


def dump_docx(path: str, out: str) -> None:
    doc = Document(path)
    lines = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            lines.append(t)
    for ti, table in enumerate(doc.tables):
        lines.append(f"\n=== TABLA {ti + 1} ===")
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " / ") for c in row.cells]
            lines.append(" | ".join(cells))
    with open(out, "w") as f:
        f.write("\n".join(lines))


def dump_xlsx(path: str, out: str) -> None:
    wb = openpyxl.load_workbook(path, data_only=True)
    lines = []
    for ws in wb.worksheets:
        lines.append(f"===== HOJA: {ws.title} ({ws.max_row}x{ws.max_column}) =====")
        for row in ws.iter_rows():
            vals = ["" if c.value is None else str(c.value).replace("\n", " / ").strip() for c in row]
            line = " | ".join(vals).rstrip(" |")
            if line.strip():
                lines.append(line)
        notes = [
            f"{c.coordinate}: {c.comment.text.strip()}"
            for row in ws.iter_rows() for c in row if c.comment
        ]
        if notes:
            lines.append("--- comentarios de celda ---")
            lines.extend(notes)
    with open(out, "w") as f:
        f.write("\n".join(lines))


def main() -> None:
    if len(sys.argv) != 2:
        sys.exit("Uso: python3 scripts/extraer_caso.py <carpeta-del-caso>")
    caso = os.path.abspath(sys.argv[1])
    if not os.path.isdir(caso):
        sys.exit(f"No existe la carpeta: {caso}")

    docx_todos = [f for f in glob.glob(os.path.join(caso, "*.docx")) if not os.path.basename(f).startswith("~$")]
    # El docx de universo se distingue por nombre de archivo (contiene
    # "universo", sin distinguir mayúsculas) — así puede convivir con el docx
    # del perfil sin que uno pise al otro.
    docx_universo = [f for f in docx_todos if "universo" in os.path.basename(f).lower()]
    docx_perfil = [f for f in docx_todos if f not in docx_universo]
    xlsx = [f for f in glob.glob(os.path.join(caso, "*.xlsx")) if not os.path.basename(f).startswith("~$")]
    pdfs = glob.glob(os.path.join(caso, "*.pdf"))
    fotos = [f for f in glob.glob(os.path.join(caso, "*.png")) + glob.glob(os.path.join(caso, "*.jpg"))
             + glob.glob(os.path.join(caso, "*.jpeg"))]
    if not docx_perfil and not xlsx:
        sys.exit(f"La carpeta no tiene ni .docx (perfil) ni .xlsx (screening): {caso}")

    out = os.path.join(caso, "_extraccion")
    os.makedirs(out, exist_ok=True)

    for f in docx_perfil:
        dump_docx(f, os.path.join(out, "perfil.txt"))
        print(f"perfil.txt     <- {os.path.basename(f)}")
    for f in docx_universo:
        dump_docx(f, os.path.join(out, "universo.txt"))
        print(f"universo.txt   <- {os.path.basename(f)}")
    for f in xlsx:
        dump_xlsx(f, os.path.join(out, "screening.txt"))
        print(f"screening.txt  <- {os.path.basename(f)}")
    if len(docx_perfil) > 1 or len(docx_universo) > 1 or len(xlsx) > 1:
        print("OJO: hay más de un docx de perfil, de universo o de xlsx; se extrajo el último de cada tipo. Revisa cuál corresponde.")
    if not docx_perfil:
        print(
            "OJO: no hay .docx de perfil en esta carpeta. Si el perfil de cargo vino en "
            "PDF, es uno de los PDF de abajo (no un CV) — identifícalo y léelo directo "
            "con el Read tool, igual que un CV; este script no lo convierte."
        )
    print(f"{len(pdfs)} PDF en la carpeta (léelos directo con el Read tool: los CV, y el perfil si vino en ese formato)")
    if fotos:
        print(f"{len(fotos)} foto(s) de candidatos en la carpeta (png/jpg) — ver SKILL.md de generar-deck sobre cómo usarlas.")
    print(f"listo → {out}")


if __name__ == "__main__":
    main()
